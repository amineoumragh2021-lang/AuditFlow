from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Titre
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "AuditPro PFA"
    subtitle.text = "Application de Gestion des Audits Industriels\nProjet Fin d'Année"
    
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    
    # Slide 2: Introduction
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Introduction"
    text_frame = content.text_frame
    text_frame.text = "AuditPro PFA est une application web Django complète pour la gestion des audits industriels"
    
    p = text_frame.add_paragraph()
    p.text = "Développée avec une interface moderne et responsive"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "Facilite la planification, l'exécution et le suivi des audits"
    p.level = 1
    
    # Slide 3: Objectifs
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Objectifs du Projet"
    text_frame = content.text_frame
    
    points = [
        "Digitaliser le processus d'audit industriel",
        "Gérer les rôles utilisateurs (Responsable, Auditeur, Client)",
        "Automatiser les checklists de conformité",
        "Générer des rapports avec scores de conformité",
        "Offrir une interface utilisateur moderne et intuitive"
    ]
    
    for point in points:
        p = text_frame.add_paragraph()
        p.text = f"• {point}"
        p.level = 0
    
    # Slide 4: Architecture Technique
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Architecture Technique"
    text_frame = content.text_frame
    
    tech_points = [
        "Backend: Django 5.0.7 (Framework Python)",
        "Base de données: MySQL via XAMPP",
        "Frontend: TailwindCSS avec animations avancées",
        "Design: Glassmorphism, responsive mobile",
        "Authentification: Système Django intégré"
    ]
    
    for point in tech_points:
        p = text_frame.add_paragraph()
        p.text = f"• {point}"
        p.level = 0
    
    # Slide 5: Rôles Utilisateurs
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Rôles Utilisateurs"
    text_frame = content.text_frame
    
    roles = [
        "Responsable Audit: Crée et planifie les audits",
        "Auditeur: Exécute les audits et remplit les checklists",
        "Client: Consulte les rapports et résultats"
    ]
    
    for role in roles:
        p = text_frame.add_paragraph()
        p.text = f"• {role}"
        p.level = 0
    
    # Slide 6: Fonctionnalités Principales
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Fonctionnalités Principales"
    text_frame = content.text_frame
    
    features = [
        "Inscription et authentification des utilisateurs",
        "Planification des audits avec dates et descriptions",
        "Affectation des auditeurs et clients aux audits",
        "Checklists par type d'audit (180 points de démonstration)",
        "Types: Qualité, Sécurité, Environnement, Interne",
        "Génération de rapports avec score de conformité",
        "Dashboard avec statistiques et compteurs animés"
    ]
    
    for feature in features:
        p = text_frame.add_paragraph()
        p.text = f"• {feature}"
        p.level = 0
    
    # Slide 7: Modèle de Données
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Modèle de Données"
    text_frame = content.text_frame
    
    models = [
        "User: Utilisateurs du système",
        "Profile: Rôles et informations complémentaires",
        "Audit: Informations sur les audits planifiés",
        "ChecklistItem: Points de contrôle par type d'audit",
        "ChecklistResponse: Réponses aux checklists",
        "Rapport: Rapports générés avec scores"
    ]
    
    for model in models:
        p = text_frame.add_paragraph()
        p.text = f"• {model}"
        p.level = 0
    
    # Slide 8: Interface Utilisateur
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Interface Utilisateur"
    text_frame = content.text_frame
    
    ui_features = [
        "Design moderne avec TailwindCSS",
        "Effets glassmorphism sur les cartes",
        "Animations fluides et transitions",
        "Responsive design (mobile, tablette, desktop)",
        "Compteurs animés pour les statistiques",
        "Navigation intuitive et ergonomique"
    ]
    
    for feature in ui_features:
        p = text_frame.add_paragraph()
        p.text = f"• {feature}"
        p.level = 0
    
    # Slide 9: Installation
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Installation"
    text_frame = content.text_frame
    
    steps = [
        "Démarrer MySQL dans XAMPP",
        "Créer la base de données audit_pfa_db",
        "Installer les dépendances: pip install -r requirements.txt",
        "Exécuter les migrations: python manage.py migrate",
        "Peupler la base: python manage.py seed_demo",
        "Lancer le serveur: python manage.py runserver"
    ]
    
    for i, step in enumerate(steps, 1):
        p = text_frame.add_paragraph()
        p.text = f"{i}. {step}"
        p.level = 0
    
    # Slide 10: Comptes de Test
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Comptes de Test"
    text_frame = content.text_frame
    
    accounts = [
        "Responsable Audit: responsable@audit.com / 123456",
        "Auditeur: auditeur@audit.com / 123456",
        "Client: client@audit.com / 123456"
    ]
    
    for account in accounts:
        p = text_frame.add_paragraph()
        p.text = f"• {account}"
        p.level = 0
    
    # Slide 11: Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = "Conclusion"
    text_frame = content.text_frame
    
    text_frame.text = "AuditPro PFA offre une solution complète et moderne pour la gestion des audits industriels"
    
    p = text_frame.add_paragraph()
    p.text = "Interface intuitive et fonctionnalités avancées"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "Architecture robuste et extensible"
    p.level = 1
    
    p = text_frame.add_paragraph()
    p.text = "Prêt pour le déploiement en production"
    p.level = 1
    
    # Save the presentation
    output_path = r"c:\Users\amine\Desktop\AuditFlow_CORRIGE_Date_Rapport_Modifier\audit_pfa2004_corrige_sans_statut\AuditPro_PFA_Presentation.pptx"
    prs.save(output_path)
    print(f"Présentation générée avec succès: {output_path}")

if __name__ == "__main__":
    create_presentation()
